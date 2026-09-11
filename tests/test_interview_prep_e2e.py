"""End-to-end CLI-level coverage for interview_prep.py.

Exercises the real `main()`/argument-parsing/file-reading path (not the
Python function directly): one run with every optional flag supplied at
once, and one run with none, asserting the zero-flag doc contains no
fabricated placeholder text anywhere.
"""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from project_1_application_engine.skills.interview_prep.scripts.interview_prep import (
    main as interview_prep_main,
)

from shared.markdown_sources import parse_star_bank
from shared.models import (
    Contact,
    ProfileBundle,
    Resume,
    SkillCategory,
    StarStory,
    WorkExperience,
)

JD_TEXT = """\
# Senior .NET Developer

**Company:** Acme

## Required Skills
- C#
- Kubernetes

## Nice to Have
- Python
"""

STAR_BANK_TEXT = """\
# STAR Story Bank

## Story 1 — SSIS Performance Improvement (Riverside Pharmacy)

**Use for:** Setback or challenge · SQL / data work

**JD tags:** `c#` `legacy-systems`

**last_used:** April 2026
**used_for:** Actalent — Software Developer

**Situation:**
Situation text.

**Task:**
Task text.

**Action:**
Action text.

**Result:**
Result text.

---
"""


def _write_jd(tmp_path: Path) -> Path:
    jd_path = tmp_path / "jd.md"
    jd_path.write_text(JD_TEXT, encoding="utf-8")
    return jd_path


def _write_star_bank(tmp_path: Path) -> Path:
    path = tmp_path / "star-bank.md"
    path.write_text(STAR_BANK_TEXT, encoding="utf-8")
    return path


def _make_bundle() -> ProfileBundle:
    resume = Resume(
        headline="Full-Stack Engineer",
        summary="Test",
        technical_skills=[SkillCategory(category="Languages", skills=["Python"])],
        experience=[
            WorkExperience(
                title="Developer",
                company="Acme Corp",
                location="",
                dates="2023-Present",
                bullets=["Built C# APIs"],
            )
        ],
        star_stories=[
            StarStory(
                title="SSIS Performance Improvement (Riverside Pharmacy)",
                jd_tags=["c#"],
                situation="Situation text.",
                task="Task text.",
                action="Action text.",
                result="Result text.",
            )
        ],
    )
    return ProfileBundle(
        bundle_version=1,
        generated=date(2026, 8, 20),
        sources={},
        contact=Contact(name="Sarah Ashford", email="andrea@example.com", phone=""),
        headline="Full-Stack Engineer",
        summary="Test",
        cert_registry=[],
        differentiators=[],
        known_genuine_gaps=[],
        resolved_framing_gaps=[],
        portfolio_projects=[],
        ats_keywords={},
        ai200_domain_coverage=[],
        resume=resume,
    )


def _write_bundle_dir(tmp_path: Path) -> Path:
    """Write a minimal profile-bundle.json + presence-bundle.json + README.md
    under tmp_path so `--bundle tmp_path` satisfies `load_profile_bundle()`."""
    bundle = _make_bundle()
    outputs = tmp_path / "outputs"
    outputs.mkdir()
    (outputs / "profile-bundle.json").write_text(json.dumps(bundle.to_dict()), encoding="utf-8")
    presence = {
        "bundle_version": 1,
        "generated": "2026-08-20",
        "headline": "Full-Stack Engineer",
        "summary": "Test",
        "cert_registry": [],
        "technical_skills": {},
        "portfolio_projects": [],
        "differentiators": [],
    }
    (outputs / "presence-bundle.json").write_text(json.dumps(presence), encoding="utf-8")
    (tmp_path / "README.md").write_text("# x", encoding="utf-8")
    return tmp_path


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_every_optional_flag_together_end_to_end(tmp_path: Path) -> None:
    bundle_root = _write_bundle_dir(tmp_path)
    jd_path = _write_jd(tmp_path)
    star_bank_path = _write_star_bank(tmp_path)

    company_research_file = tmp_path / "company-research.md"
    company_research_file.write_text(
        "Acme just closed a Series B and is expanding its platform team.", encoding="utf-8"
    )
    recruiter_briefing_file = tmp_path / "recruiter-briefing.md"
    recruiter_briefing_file.write_text(
        "Panel format, three interviewers, focus on system design.", encoding="utf-8"
    )
    qa_cards_file = tmp_path / "qa-cards.json"
    qa_cards_file.write_text(
        json.dumps(
            [
                {
                    "question": "Walk me through your resume.",
                    "answer": (
                        "I started as a support engineer and moved into full-stack development."
                    ),
                    "tip": "Keep it under two minutes.",
                    "category": "resume_walkthrough",
                }
            ]
        ),
        encoding="utf-8",
    )
    salary_section_file = tmp_path / "salary.md"
    salary_section_file.write_text(
        "Target range: $120K-$140K CAD based on current market research.", encoding="utf-8"
    )
    out_file = tmp_path / "prep.docx"

    rc = interview_prep_main(
        [
            str(jd_path),
            "--stage",
            "final_round",
            "--out",
            str(out_file),
            "--bundle",
            str(bundle_root),
            "--company-research",
            str(company_research_file),
            "--recruiter-briefing",
            str(recruiter_briefing_file),
            "--qa-cards",
            str(qa_cards_file),
            "--salary-section",
            str(salary_section_file),
            "--update-star-bank",
            "--star-bank-path",
            str(star_bank_path),
        ]
    )
    assert rc == 0
    assert out_file.exists()

    text = _docx_text(out_file)
    assert "Senior .NET Developer" in text
    assert "Acme" in text
    assert "Final Round" in text
    assert "Acme just closed a Series B and is expanding its platform team." in text
    assert "Recruiter Briefing" in text
    assert "Panel format, three interviewers, focus on system design." in text
    assert "Strong Match" in text
    assert "Walk me through your resume." in text
    assert "SSIS Performance Improvement" in text
    assert "Target range: $120K-$140K CAD based on current market research." in text

    updated = parse_star_bank(star_bank_path.read_text(encoding="utf-8"))
    assert updated[0]["last_used"] == date.today().strftime("%B %Y")
    assert "Acme — Senior .NET Developer" in updated[0]["used_for"]

    sidecar_path = out_file.parent / (out_file.name + ".meta.json")
    assert sidecar_path.exists()
    meta = json.loads(sidecar_path.read_text(encoding="utf-8"))
    assert meta["script"] == "interview_prep.py"
    assert meta["fit_rating"]
    assert str(jd_path) in meta["inputs"]


def test_zero_optional_flags_minimal_valid_doc_with_no_fabrication(tmp_path: Path) -> None:
    bundle_root = _write_bundle_dir(tmp_path)
    jd_path = _write_jd(tmp_path)
    out_file = tmp_path / "prep.docx"

    rc = interview_prep_main([str(jd_path), "--out", str(out_file), "--bundle", str(bundle_root)])
    assert rc == 0
    assert out_file.exists()
    assert out_file.stat().st_size > 0

    text = _docx_text(out_file)
    assert text.strip()
    assert "Technical Screen" in text  # default stage, no --stage supplied

    for placeholder in ("TBD", "N/A", "visit their website", "[placeholder]", "Lorem ipsum"):
        assert placeholder not in text

    # No hybrid content supplied — none of their sections should appear at all.
    assert "COMPANY RESEARCH" not in text
    assert "Recruiter Briefing" not in text
    assert "SALARY / NEGOTIATION" not in text

    sidecar_path = out_file.parent / (out_file.name + ".meta.json")
    assert sidecar_path.exists()
