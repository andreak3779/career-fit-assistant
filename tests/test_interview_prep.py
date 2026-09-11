"""Tests for interview_prep.py — the DOCX interview-prep generator."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from project_1_application_engine.skills.interview_prep.scripts.interview_prep import (
    generate_interview_prep,
)
from project_1_application_engine.skills.interview_prep.scripts.interview_prep import (
    main as interview_prep_main,
)

from shared.markdown_sources import parse_star_bank
from shared.models import (
    Contact,
    ProfileBundle,
    Resume,
    SkillCategory,
    StarStory,
    WorkExperience,
)

REAL_STAR_BANK_PATH = (
    Path(__file__).resolve().parents[1]
    / "project-1-application-engine"
    / "reference"
    / "star-bank.md"
)

SAMPLE_STAR_BANK_TEXT = """\
# STAR Story Bank

## Story 1 — SSIS Performance Improvement (Riverside Pharmacy)

**Use for:** Setback or challenge · SQL / data work

**JD tags:** `sql` `legacy-systems`

**last_used:** April 2026
**used_for:** Actalent — Software Developer

**Situation:**
Situation text.

**Task:**
Task text.

**Action:**
Action text.

**Result:**
Result text.

---
"""


def _write_star_bank(tmp_path: Path) -> Path:
    path = tmp_path / "star-bank.md"
    path.write_text(SAMPLE_STAR_BANK_TEXT, encoding="utf-8")
    return path


JD_TEXT = """\
# Senior .NET Developer

**Company:** Acme

## Required Skills
- C#
- Kubernetes

## Nice to Have
- Python
"""


def _make_bundle(*, star_stories: list[StarStory] | None = None) -> ProfileBundle:
    resume = Resume(
        headline="Full-Stack Engineer",
        summary="Test",
        technical_skills=[SkillCategory(category="Languages", skills=["Python"])],
        experience=[
            WorkExperience(
                title="Developer",
                company="Acme Corp",
                location="",
                dates="2023-Present",
                bullets=["Built C# APIs"],
            )
        ],
        star_stories=star_stories or [],
    )
    return ProfileBundle(
        bundle_version=1,
        generated=date(2026, 8, 20),
        sources={},
        contact=Contact(name="Sarah Ashford", email="andrea@example.com", phone=""),
        headline="Full-Stack Engineer",
        summary="Test",
        cert_registry=[],
        differentiators=[],
        known_genuine_gaps=[],
        resolved_framing_gaps=[],
        portfolio_projects=[],
        ats_keywords={},
        ai200_domain_coverage=[],
        resume=resume,
    )


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _write_jd(tmp_path: Path) -> Path:
    jd_path = tmp_path / "jd.md"
    jd_path.write_text(JD_TEXT, encoding="utf-8")
    return jd_path


GAP_COACHING_TEXT = """\
# Gap Coaching Reference — Test

## Kubernetes

**Status:** 🔴 Genuine Gap — no Kubernetes exposure

**What to say:**
"I have not worked with Kubernetes directly, but I have deployed containerized apps with Docker."

**What not to say:**
- Do not claim Kubernetes production experience

**Best STAR bridge:** Story 4 (Learning ramp)

**Honest floor:** Zero Kubernetes production experience. Docker experience present.

---

## C#

**Status:** 🟡 Framing gap — production C# experience should be surfaced confidently

**What to say:**
"C# has been my primary language across every production role."

**What not to say:** Nothing — lead with this confidently.

**Best STAR bridge:** Story 1 (SSIS performance)

**Honest floor:** Production C# experience across multiple employers.
"""


def _write_gap_coaching(tmp_path: Path) -> Path:
    path = tmp_path / "gap-coaching.md"
    path.write_text(GAP_COACHING_TEXT, encoding="utf-8")
    return path


def test_banner_shows_title_company_and_stage(tmp_path: Path) -> None:
    bundle = _make_bundle()
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path), bundle=bundle, out_path=tmp_path / "prep.docx", root=tmp_path
    )
    text = _docx_text(out_path)
    assert "Senior .NET Developer" in text
    assert "Acme" in text
    assert "Technical Screen" in text


def test_experience_match_table_has_a_row_per_required_and_nice_to_have_skill(
    tmp_path: Path,
) -> None:
    bundle = _make_bundle()
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path), bundle=bundle, out_path=tmp_path / "prep.docx", root=tmp_path
    )
    text = _docx_text(out_path)
    assert "Strong Match" in text
    assert "Acme Corp: production" in text
    assert "Genuine Gap" in text
    assert "Not present in resume, portfolio, or coursework" in text


def test_grounded_skill_gets_a_qa_card_with_evidence_answer(tmp_path: Path) -> None:
    bundle = _make_bundle()
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path), bundle=bundle, out_path=tmp_path / "prep.docx", root=tmp_path
    )
    text = _docx_text(out_path)
    assert "walk me through how you used c#" in text.lower()
    assert "Evidence on file: Acme Corp: production" in text


def test_ungrounded_skill_gets_an_honesty_reminder_not_a_fabricated_answer(
    tmp_path: Path,
) -> None:
    bundle = _make_bundle()
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path), bundle=bundle, out_path=tmp_path / "prep.docx", root=tmp_path
    )
    text = _docx_text(out_path)
    assert "if kubernetes comes up, answer honestly" in text.lower()


def test_gap_coaching_entry_replaces_honesty_reminder_with_a_real_qa_card(
    tmp_path: Path,
) -> None:
    bundle = _make_bundle()
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path),
        bundle=bundle,
        out_path=tmp_path / "prep.docx",
        root=tmp_path,
        gap_coaching_path=_write_gap_coaching(tmp_path),
    )
    text = _docx_text(out_path)
    assert "if kubernetes comes up, answer honestly" not in text.lower()
    assert "I have not worked with Kubernetes directly" in text
    assert "Avoid: Do not claim Kubernetes production experience" in text
    assert "Zero Kubernetes production experience" in text
    assert "Story 4 (Learning ramp)" in text


def test_gap_coaching_entry_overrides_grounded_evidence_stub_answer(tmp_path: Path) -> None:
    bundle = _make_bundle()
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path),
        bundle=bundle,
        out_path=tmp_path / "prep.docx",
        root=tmp_path,
        gap_coaching_path=_write_gap_coaching(tmp_path),
    )
    text = _docx_text(out_path)
    assert "C# has been my primary language across every production role." in text
    assert "Evidence on file: Acme Corp: production" not in text


def test_missing_gap_coaching_file_falls_back_cleanly(tmp_path: Path) -> None:
    bundle = _make_bundle()
    missing_path = tmp_path / "does-not-exist.md"
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path),
        bundle=bundle,
        out_path=tmp_path / "prep.docx",
        root=tmp_path,
        gap_coaching_path=missing_path,
    )
    text = _docx_text(out_path)
    assert "Evidence on file: Acme Corp: production" in text
    assert "if kubernetes comes up, answer honestly" in text.lower()


def test_star_card_count_gated_by_stage(tmp_path: Path) -> None:
    stories = [
        StarStory(
            title=f"Story {i}",
            jd_tags=[tag],
            situation="S",
            task="T",
            action="A",
            result="R",
        )
        for i, tag in enumerate(["c#", "python", "kubernetes"])
    ]
    bundle = _make_bundle(star_stories=stories)

    phone_out, _, _ = generate_interview_prep(
        _write_jd(tmp_path),
        bundle=bundle,
        out_path=tmp_path / "phone.docx",
        stage="phone_screen",
        root=tmp_path,
    )
    phone_text = _docx_text(phone_out)
    assert (
        phone_text.count("Story 0") + phone_text.count("Story 1") + phone_text.count("Story 2") == 2
    )

    final_out, _, _ = generate_interview_prep(
        _write_jd(tmp_path),
        bundle=bundle,
        out_path=tmp_path / "final.docx",
        stage="final_round",
        root=tmp_path,
    )
    final_text = _docx_text(final_out)
    assert "Story 0" in final_text
    assert "Story 1" in final_text
    assert "Story 2" in final_text


def test_star_card_shows_adapt_when_as_reference_text(tmp_path: Path) -> None:
    stories = [
        StarStory(
            title="Story A",
            jd_tags=["c#"],
            situation="S",
            task="T",
            action="A",
            result="R",
            adapt_when=["Emphasize the API work for backend-heavy roles"],
        )
    ]
    bundle = _make_bundle(star_stories=stories)
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path), bundle=bundle, out_path=tmp_path / "prep.docx", root=tmp_path
    )
    text = _docx_text(out_path)
    assert "Adapt when:" in text
    assert "Emphasize the API work for backend-heavy roles" in text


def test_no_star_stories_shows_honest_placeholder(tmp_path: Path) -> None:
    bundle = _make_bundle(star_stories=[])
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path), bundle=bundle, out_path=tmp_path / "prep.docx", root=tmp_path
    )
    text = _docx_text(out_path)
    assert "No tagged STAR stories found" in text


def test_fit_rating_section_present(tmp_path: Path) -> None:
    bundle = _make_bundle()
    out_path, _, result = generate_interview_prep(
        _write_jd(tmp_path), bundle=bundle, out_path=tmp_path / "prep.docx", root=tmp_path
    )
    text = _docx_text(out_path)
    assert result.rating in text
    assert result.rationale in text


def test_output_is_docx_not_markdown(tmp_path: Path) -> None:
    bundle = _make_bundle()
    out_path, _, _ = generate_interview_prep(_write_jd(tmp_path), bundle=bundle, root=tmp_path)
    assert out_path.suffix == ".docx"
    assert out_path.name == "acme_senior_net_developer_InterviewPrep.docx"


def test_writes_sidecar_with_fit_rating_and_evidence(tmp_path: Path) -> None:
    bundle = _make_bundle()
    out_path, sidecar_path, result = generate_interview_prep(
        _write_jd(tmp_path), bundle=bundle, out_path=tmp_path / "prep.docx", root=tmp_path
    )
    assert out_path.exists()
    assert sidecar_path.exists()
    meta = json.loads(sidecar_path.read_text(encoding="utf-8"))
    assert meta["script"] == "interview_prep.py"
    assert meta["fit_rating"] == result.rating


def test_main_cli_end_to_end_with_stage_flag(tmp_path: Path) -> None:
    jd_path = _write_jd(tmp_path)
    out_file = tmp_path / "prep.docx"
    bundle = _make_bundle()

    # Write a minimal profile-bundle.json + presence-bundle.json for load_profile_bundle.
    outputs = tmp_path / "outputs"
    outputs.mkdir()
    (outputs / "profile-bundle.json").write_text(json.dumps(bundle.to_dict()), encoding="utf-8")
    presence = {
        "bundle_version": 1,
        "generated": "2026-08-20",
        "headline": "Full-Stack Engineer",
        "summary": "Test",
        "cert_registry": [],
        "technical_skills": {},
        "portfolio_projects": [],
        "differentiators": [],
    }
    (outputs / "presence-bundle.json").write_text(json.dumps(presence), encoding="utf-8")
    (tmp_path / "README.md").write_text("# x")

    rc = interview_prep_main(
        [
            str(jd_path),
            "--stage",
            "phone_screen",
            "--out",
            str(out_file),
            "--bundle",
            str(tmp_path),
        ]
    )
    assert rc == 0
    assert out_file.exists()


def test_write_star_bank_true_updates_selected_story(tmp_path: Path) -> None:
    star_bank_path = _write_star_bank(tmp_path)
    stories = [
        StarStory(
            title="SSIS Performance Improvement (Riverside Pharmacy)",
            jd_tags=["c#"],
            situation="S",
            task="T",
            action="A",
            result="R",
        )
    ]
    bundle = _make_bundle(star_stories=stories)
    generate_interview_prep(
        _write_jd(tmp_path),
        bundle=bundle,
        out_path=tmp_path / "prep.docx",
        root=tmp_path,
        write_star_bank=True,
        star_bank_path=star_bank_path,
    )
    updated = parse_star_bank(star_bank_path.read_text(encoding="utf-8"))
    assert updated[0]["last_used"] == date.today().strftime("%B %Y")
    assert updated[0]["used_for"] == [
        "Actalent — Software Developer",
        "Acme — Senior .NET Developer",
    ]


def test_write_star_bank_false_leaves_file_untouched(tmp_path: Path) -> None:
    star_bank_path = _write_star_bank(tmp_path)
    original_text = star_bank_path.read_text(encoding="utf-8")
    stories = [
        StarStory(
            title="SSIS Performance Improvement (Riverside Pharmacy)",
            jd_tags=["c#"],
            situation="S",
            task="T",
            action="A",
            result="R",
        )
    ]
    bundle = _make_bundle(star_stories=stories)
    generate_interview_prep(
        _write_jd(tmp_path),
        bundle=bundle,
        out_path=tmp_path / "prep.docx",
        root=tmp_path,
        star_bank_path=star_bank_path,
    )
    assert star_bank_path.read_text(encoding="utf-8") == original_text


def test_real_star_bank_reference_file_mtime_unchanged_by_star_bank_path_override(
    tmp_path: Path,
) -> None:
    before_mtime = REAL_STAR_BANK_PATH.stat().st_mtime
    star_bank_path = _write_star_bank(tmp_path)
    stories = [
        StarStory(
            title="SSIS Performance Improvement (Riverside Pharmacy)",
            jd_tags=["c#"],
            situation="S",
            task="T",
            action="A",
            result="R",
        )
    ]
    bundle = _make_bundle(star_stories=stories)
    generate_interview_prep(
        _write_jd(tmp_path),
        bundle=bundle,
        out_path=tmp_path / "prep.docx",
        root=tmp_path,
        write_star_bank=True,
        star_bank_path=star_bank_path,
    )
    assert REAL_STAR_BANK_PATH.stat().st_mtime == before_mtime


def test_company_research_renders_verbatim_when_given(tmp_path: Path) -> None:
    bundle = _make_bundle()
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path),
        bundle=bundle,
        out_path=tmp_path / "prep.docx",
        root=tmp_path,
        company_research="Acme just raised a Series C and is expanding the platform team.",
    )
    text = _docx_text(out_path)
    assert "COMPANY RESEARCH" in text
    assert "Acme just raised a Series C and is expanding the platform team." in text


def test_company_research_section_omitted_when_absent(tmp_path: Path) -> None:
    bundle = _make_bundle()
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path), bundle=bundle, out_path=tmp_path / "prep.docx", root=tmp_path
    )
    text = _docx_text(out_path)
    assert "COMPANY RESEARCH" not in text


def test_company_research_section_omitted_when_blank(tmp_path: Path) -> None:
    bundle = _make_bundle()
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path),
        bundle=bundle,
        out_path=tmp_path / "prep.docx",
        root=tmp_path,
        company_research="   \n  ",
    )
    text = _docx_text(out_path)
    assert "COMPANY RESEARCH" not in text


def test_recruiter_briefing_renders_as_callout_when_given(tmp_path: Path) -> None:
    bundle = _make_bundle()
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path),
        bundle=bundle,
        out_path=tmp_path / "prep.docx",
        root=tmp_path,
        recruiter_briefing="Panel format, three interviewers, focus on system design.",
    )
    text = _docx_text(out_path)
    assert "Recruiter Briefing" in text
    assert "Panel format, three interviewers, focus on system design." in text


def test_recruiter_briefing_omitted_when_absent(tmp_path: Path) -> None:
    bundle = _make_bundle()
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path), bundle=bundle, out_path=tmp_path / "prep.docx", root=tmp_path
    )
    text = _docx_text(out_path)
    assert "Recruiter Briefing" not in text


def test_qa_cards_appear_after_scripts_own_grounded_cards(tmp_path: Path) -> None:
    bundle = _make_bundle()
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path),
        bundle=bundle,
        out_path=tmp_path / "prep.docx",
        root=tmp_path,
        qa_cards=[
            {
                "question": "Walk me through your resume.",
                "answer": "I started as a support engineer and moved into full-stack development.",
                "category": "resume_walkthrough",
            }
        ],
    )
    doc = Document(str(out_path))
    paragraph_texts = [p.text for p in doc.paragraphs]
    grounded_idx = next(
        i for i, t in enumerate(paragraph_texts) if "walk me through how you used c#" in t.lower()
    )
    supplied_idx = next(
        i for i, t in enumerate(paragraph_texts) if t == "Walk me through your resume."
    )
    assert supplied_idx > grounded_idx


def test_qa_cards_none_or_empty_renders_only_scripts_own_cards(tmp_path: Path) -> None:
    bundle = _make_bundle()
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path),
        bundle=bundle,
        out_path=tmp_path / "prep.docx",
        root=tmp_path,
        qa_cards=[],
    )
    text = _docx_text(out_path)
    assert "walk me through how you used c#" in text.lower()


def test_qa_cards_entry_missing_answer_raises_value_error(tmp_path: Path) -> None:
    bundle = _make_bundle()
    try:
        generate_interview_prep(
            _write_jd(tmp_path),
            bundle=bundle,
            out_path=tmp_path / "prep.docx",
            root=tmp_path,
            qa_cards=[{"question": "What is your greatest weakness?"}],
        )
        raise AssertionError("expected ValueError")
    except ValueError as exc:
        assert "answer" in str(exc)


def test_qa_cards_entry_missing_question_raises_value_error(tmp_path: Path) -> None:
    bundle = _make_bundle()
    try:
        generate_interview_prep(
            _write_jd(tmp_path),
            bundle=bundle,
            out_path=tmp_path / "prep.docx",
            root=tmp_path,
            qa_cards=[{"answer": "I'd say attention to detail."}],
        )
        raise AssertionError("expected ValueError")
    except ValueError as exc:
        assert "question" in str(exc)


def test_qa_cards_validation_error_prevents_file_write(tmp_path: Path) -> None:
    bundle = _make_bundle()
    out_path = tmp_path / "prep.docx"
    try:
        generate_interview_prep(
            _write_jd(tmp_path),
            bundle=bundle,
            out_path=out_path,
            root=tmp_path,
            qa_cards=[{"question": "Missing an answer"}],
        )
    except ValueError:
        pass
    assert not out_path.exists()


def test_salary_section_renders_for_final_round(tmp_path: Path) -> None:
    bundle = _make_bundle()
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path),
        bundle=bundle,
        out_path=tmp_path / "prep.docx",
        root=tmp_path,
        stage="final_round",
        salary_section="Target range: $120K-$140K CAD based on market research.",
    )
    text = _docx_text(out_path)
    assert "SALARY / NEGOTIATION" in text
    assert "Target range: $120K-$140K CAD based on market research." in text


def test_salary_section_omitted_when_stage_not_final_round(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    bundle = _make_bundle()
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path),
        bundle=bundle,
        out_path=tmp_path / "prep.docx",
        root=tmp_path,
        stage="technical_screen",
        salary_section="Target range: $120K-$140K CAD based on market research.",
    )
    text = _docx_text(out_path)
    assert "SALARY / NEGOTIATION" not in text
    assert "Target range" not in text
    captured = capsys.readouterr()
    assert "Salary section supplied but stage is not final_round — omitted." in captured.out


def test_salary_section_omitted_when_not_given_even_at_final_round(tmp_path: Path) -> None:
    bundle = _make_bundle()
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path),
        bundle=bundle,
        out_path=tmp_path / "prep.docx",
        root=tmp_path,
        stage="final_round",
    )
    text = _docx_text(out_path)
    assert "SALARY / NEGOTIATION" not in text


def test_all_four_hybrid_params_together_end_to_end(tmp_path: Path) -> None:
    bundle = _make_bundle()
    out_path, _, _ = generate_interview_prep(
        _write_jd(tmp_path),
        bundle=bundle,
        out_path=tmp_path / "prep.docx",
        root=tmp_path,
        stage="final_round",
        company_research="Acme is scaling its platform team after a Series C.",
        recruiter_briefing="Panel format, focus on system design.",
        qa_cards=[
            {
                "question": "Walk me through your resume.",
                "answer": "Support engineer to full-stack developer.",
            }
        ],
        salary_section="Target range: $120K-$140K CAD.",
    )
    text = _docx_text(out_path)
    assert "COMPANY RESEARCH" in text
    assert "Acme is scaling its platform team after a Series C." in text
    assert "Recruiter Briefing" in text
    assert "Panel format, focus on system design." in text
    assert "Walk me through your resume." in text
    assert "SALARY / NEGOTIATION" in text
    assert "Target range: $120K-$140K CAD." in text
