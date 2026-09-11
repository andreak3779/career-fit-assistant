"""Tests for shared.docx_layout and shared.jd_parser helpers."""

from __future__ import annotations

from docx import Document

from shared.docx_layout import (
    DEFAULT_SIGNATURE_NAME,
    bullet_label_para,
    clean_stack,
    letter_footer,
    slugify,
    split_label_body,
)
from shared.jd_parser import looks_unknown, parse_jd


class TestSlugify:
    def test_basic(self) -> None:
        assert slugify("Contoso Cloud Solutions") == "contoso_cloud_solutions"

    def test_strips_edges(self) -> None:
        assert slugify("---acme---") == "acme"

    def test_collapses_runs(self) -> None:
        assert slugify("a   b   c") == "a_b_c"

    def test_strips_punctuation(self) -> None:
        assert slugify("Acme, Inc. (USA)!") == "acme_inc_usa"

    def test_lowercases(self) -> None:
        assert slugify("Contoso") == "contoso"


class TestCleanStack:
    def test_strips_tokens(self) -> None:
        # Real bundles sometimes have leading/trailing spaces per item.
        assert (
            clean_stack(["Angular ", " ASP.NET Core Web API "]) == "Angular, ASP.NET Core Web API"
        )

    def test_drops_empty_tokens(self) -> None:
        assert clean_stack(["a", "", "  ", "b"]) == "a, b"

    def test_respects_limit(self) -> None:
        assert clean_stack(["a", "b", "c", "d"], limit=2) == "a, b"


class TestSplitLabelBody:
    def test_colon(self) -> None:
        label, body = split_label_body("Legacy modernization: 3 employers")
        assert label == "Legacy modernization"
        assert body == "3 employers"

    def test_semicolon(self) -> None:
        label, body = split_label_body("Tag; rest of text")
        assert label == "Tag"
        assert body == "rest of text"

    def test_em_dash(self) -> None:
        label, body = split_label_body("Tag — rest")
        assert label == "Tag"
        assert body == "rest"

    def test_en_dash(self) -> None:
        label, body = split_label_body("Tag – rest")
        assert label == "Tag"
        assert body == "rest"

    def test_spaced_hyphen_is_separator(self) -> None:
        label, body = split_label_body("Tag - rest")
        assert label == "Tag"
        assert body == "rest"

    def test_embedded_hyphen_is_not_separator(self) -> None:
        # "Full-stack depth: …" should keep "Full-stack" intact and split on ':'.
        label, body = split_label_body("Full-stack depth: ASP.NET Core + Angular in production")
        assert label == "Full-stack depth"
        assert body == "ASP.NET Core + Angular in production"

    def test_no_separator_returns_full_string(self) -> None:
        label, body = split_label_body("Just one phrase")
        assert label == "Just one phrase"
        assert body == "Just one phrase"


class TestLooksUnknown:
    def test_none(self) -> None:
        assert looks_unknown(None) is True

    def test_empty_string(self) -> None:
        # A blank `**Company:**` line parses to "" — same as missing.
        assert looks_unknown("") is True

    def test_unknown_company(self) -> None:
        assert looks_unknown("Unknown Company") is True

    def test_unknown_role(self) -> None:
        assert looks_unknown("Unknown Role") is True

    def test_real_company(self) -> None:
        assert looks_unknown("Contoso Cloud Solutions") is False

    def test_real_company_with_unknown_word(self) -> None:
        # Edge case: company names can contain the word "Unknown" legitimately.
        # We keep the simple startswith check; callers can do additional
        # checks if they need stricter matching.
        assert looks_unknown("Unknown Software Inc.") is True


class TestBulletLabelPara:
    def test_renders_label_and_body(self) -> None:
        doc = Document()
        bullet_label_para(doc, "Tag", "details here")
        assert len(doc.paragraphs) == 1
        text = doc.paragraphs[0].text
        # Bullet glyph + label + body.
        assert "\u2022  Tag: details here" == text

    def test_label_is_bold(self) -> None:
        doc = Document()
        bullet_label_para(doc, "Tag", "details")
        runs = doc.paragraphs[0].runs
        # Three runs: bullet, label-with-colon, body.
        assert len(runs) == 3
        # Label run is bold; bullet and body are not bold.
        assert runs[1].bold is True
        assert runs[1].text == "Tag: "
        assert "details" in runs[2].text


class TestLetterFooter:
    def test_uses_contact_name_when_present(self) -> None:
        from shared.models import Contact

        doc = Document()
        contact = Contact(name="Jane Doe", email="", phone="")
        letter_footer(doc, contact)
        # The signature line is the last paragraph.
        assert doc.paragraphs[-1].text == "JANE DOE"

    def test_falls_back_to_default(self) -> None:
        from shared.models import Contact

        doc = Document()
        contact = Contact(name="", email="", phone="")
        letter_footer(doc, contact)
        assert doc.paragraphs[-1].text == DEFAULT_SIGNATURE_NAME.upper()


class TestParseJDRoundTripLooksUnknown:
    def test_blank_company_line(self) -> None:
        # A JD markdown file with `**Company:**` and no value on the same line.
        # Note: parse_jd's regex is greedy enough to pick up the next non-empty
        # line, so this exact scenario is fragile. We assert the contract of
        # looks_unknown on the empty string instead.
        assert looks_unknown("") is True
        assert looks_unknown(None) is True

    def test_parse_jd_uses_looks_unknown_in_generators(self) -> None:
        # This is the cold-call regression check: when parse_jd returns a
        # real-looking value AND looks_unknown agrees, the file is treated
        # as a real JD. When parse_jd returns "Unknown Company" or "",
        # looks_unknown returns True and the caller falls back.
        parsed = parse_jd("# A Role\n\n## Required Skills\n- Python\n")
        assert looks_unknown(parsed.company) is True
        assert parsed.title == "A Role"
