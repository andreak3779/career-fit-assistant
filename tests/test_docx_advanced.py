"""Tests for shared.docx_advanced — DOCX shading/border/table/card primitives.

python-docx has no high-level API for background fill or the callout-box/
card visual language interview_prep.py needs, so these helpers inject raw
oxml (`<w:shd>`, `<w:pBdr>`) directly. Assertions check the actual XML
produced, not just "it didn't crash" — that's the only way to verify oxml
correctness without opening the file in Word/LibreOffice.
"""

from __future__ import annotations

from docx import Document

from shared.docx_advanced import (
    AMBER_HEX,
    AMBERB_HEX,
    BLUE_HEX,
    GREEN_HEX,
    LGRAY_HEX,
    NAVY_HEX,
    RED_HEX,
    banner,
    callout_box,
    qa_card,
    set_paragraph_left_border,
    set_paragraph_top_border,
    shade_cell,
    shade_paragraph,
    star_card,
    two_column_table,
)


def _doc():
    return Document()


class TestShadeCell:
    def test_injects_shd_with_correct_fill(self) -> None:
        doc = _doc()
        table = doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        shade_cell(cell, NAVY_HEX)
        xml = cell._tc.xml
        assert "w:shd" in xml
        assert NAVY_HEX in xml


class TestShadeParagraph:
    def test_injects_shd_with_correct_fill(self) -> None:
        doc = _doc()
        p = doc.add_paragraph("text")
        shade_paragraph(p, AMBER_HEX)
        xml = p._p.xml
        assert "w:shd" in xml
        assert AMBER_HEX in xml


class TestParagraphBorders:
    def test_left_border_has_correct_side_and_color(self) -> None:
        doc = _doc()
        p = doc.add_paragraph("text")
        set_paragraph_left_border(p, AMBERB_HEX)
        xml = p._p.xml
        assert "w:pBdr" in xml
        assert "w:left" in xml
        assert AMBERB_HEX in xml

    def test_top_border_has_correct_side_and_color(self) -> None:
        doc = _doc()
        p = doc.add_paragraph("text")
        set_paragraph_top_border(p, BLUE_HEX)
        xml = p._p.xml
        assert "w:pBdr" in xml
        assert "w:top" in xml
        assert BLUE_HEX in xml

    def test_left_and_top_can_coexist_on_same_paragraph(self) -> None:
        doc = _doc()
        p = doc.add_paragraph("text")
        set_paragraph_left_border(p, AMBERB_HEX)
        set_paragraph_top_border(p, BLUE_HEX)
        xml = p._p.xml
        assert xml.count("w:pBdr") == 2  # one open + one close tag, single element
        assert "w:left" in xml
        assert "w:top" in xml


class TestTwoColumnTable:
    def test_structure_and_content(self) -> None:
        doc = _doc()
        table = two_column_table(
            doc,
            ["They Need", "Your Evidence"],
            [("C#", "Fieldstone Benefits Administrators, production"), ("Azure", "coursework")],
        )
        assert len(table.rows) == 3  # 1 header + 2 data rows
        assert len(table.columns) == 2
        assert table.cell(0, 0).text == "They Need"
        assert table.cell(0, 1).text == "Your Evidence"
        assert table.cell(1, 0).text == "C#"
        assert table.cell(2, 1).text == "coursework"

    def test_header_row_is_shaded(self) -> None:
        doc = _doc()
        table = two_column_table(doc, ["A", "B"], [("x", "y")])
        assert NAVY_HEX in table.cell(0, 0)._tc.xml

    def test_alternating_rows_shaded(self) -> None:
        doc = _doc()
        table = two_column_table(doc, ["A", "B"], [("r1c1", "r1c2"), ("r2c1", "r2c2")])
        # Row 1 (odd data row) unshaded, row 2 (even data row) shaded LGRAY.
        assert LGRAY_HEX not in table.cell(1, 0)._tc.xml
        assert LGRAY_HEX in table.cell(2, 0)._tc.xml


class TestBanner:
    def test_builds_one_by_one_table_with_title_and_subtitles(self) -> None:
        doc = _doc()
        table = banner(doc, title="Senior .NET Developer", subtitle_lines=["Acme Corp", "Aug 2026"])
        assert len(table.rows) == 1
        assert len(table.columns) == 1
        cell_text = table.cell(0, 0).text
        assert "Senior .NET Developer" in cell_text
        assert "Acme Corp" in cell_text
        assert "Aug 2026" in cell_text
        assert NAVY_HEX in table.cell(0, 0)._tc.xml

    def test_no_subtitles_still_works(self) -> None:
        doc = _doc()
        table = banner(doc, title="Just a title")
        assert table.cell(0, 0).text == "Just a title"


class TestCalloutBox:
    def test_label_and_body_present_with_fill_and_border(self) -> None:
        doc = _doc()
        p = callout_box(doc, "Gap note", "Course-level only.")
        assert p.text == "Gap note: Course-level only."
        assert AMBER_HEX in p._p.xml
        assert AMBERB_HEX in p._p.xml


class TestQACard:
    def test_question_answer_tip_and_gap_note_all_render(self) -> None:
        doc = _doc()
        qa_card(
            doc,
            "Walk me through your background",
            "I have 10+ years...",
            tip="Keep it under 90 seconds.",
            gap_note="Course-level only.",
        )
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Walk me through your background" in text
        assert "I have 10+ years..." in text
        assert "Tip: Keep it under 90 seconds." in text
        assert "Gap note: Course-level only." in text

    def test_tip_and_gap_note_omitted_when_not_given(self) -> None:
        doc = _doc()
        qa_card(doc, "Q", "A")
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Tip:" not in text
        assert "Gap note:" not in text

    def test_question_paragraph_has_blue_top_border(self) -> None:
        doc = _doc()
        qa_card(doc, "Q", "A")
        question_para = doc.paragraphs[0]
        assert "w:top" in question_para._p.xml
        assert BLUE_HEX in question_para._p.xml


class TestStarCard:
    def test_all_star_rows_present_with_correct_colors(self) -> None:
        doc = _doc()
        star_card(doc, "SSIS Improvement", ["sql", "performance"], "Sit.", "Task.", "Act.", "Res.")
        paras = doc.paragraphs
        text = "\n".join(p.text for p in paras)
        assert "SSIS Improvement" in text
        assert "sql, performance" in text
        assert "S: Sit." in text
        assert "T: Task." in text
        assert "A: Act." in text
        assert "R: Res." in text
        # Find each row and confirm its own color, not just present somewhere.
        s_row = next(p for p in paras if p.text.startswith("S:"))
        t_row = next(p for p in paras if p.text.startswith("T:"))
        a_row = next(p for p in paras if p.text.startswith("A:"))
        r_row = next(p for p in paras if p.text.startswith("R:"))
        assert BLUE_HEX in s_row._p.xml
        assert GREEN_HEX in t_row._p.xml
        assert AMBER_HEX in a_row._p.xml
        assert RED_HEX in r_row._p.xml

    def test_adapt_when_rendered_as_reference_text_not_rewritten(self) -> None:
        doc = _doc()
        star_card(
            doc,
            "Story",
            [],
            "S",
            "T",
            "A",
            "R",
            adapt_when=["Emphasize Python if data-focused JD", "Lead with SQL for DB roles"],
        )
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Adapt when:" in text
        assert "Emphasize Python if data-focused JD" in text
        assert "Lead with SQL for DB roles" in text

    def test_adapt_when_omitted_when_not_given(self) -> None:
        doc = _doc()
        star_card(doc, "Story", [], "S", "T", "A", "R")
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Adapt when:" not in text

    def test_header_shows_title_without_tags_when_no_tags(self) -> None:
        doc = _doc()
        star_card(doc, "Story Title", [], "S", "T", "A", "R")
        assert doc.paragraphs[0].text == "Story Title"


def test_all_primitives_compose_into_one_valid_document(tmp_path) -> None:
    """Round-trip smoke test: every primitive in one doc, saved and reopened."""
    doc = _doc()
    banner(doc, title="Senior .NET Developer", subtitle_lines=["Acme Corp"])
    two_column_table(doc, ["They Need", "Your Evidence"], [("C#", "production")])
    qa_card(doc, "Q", "A", tip="A tip.", gap_note="A gap note.")
    star_card(doc, "Story", ["sql"], "S", "T", "A", "R", adapt_when=["Adapt this."])
    callout_box(doc, "Note", "Standalone.")

    out = tmp_path / "smoke.docx"
    doc.save(str(out))
    assert out.exists() and out.stat().st_size > 0

    reopened = Document(str(out))
    assert len(reopened.tables) == 2  # banner + two_column_table
