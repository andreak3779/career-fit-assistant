"""Richer DOCX building blocks: shaded/bordered boxes, tables, banners, cards.

``shared/docx_layout.py`` only ever emits plain single-run paragraphs — no
script in this repo has ever built a python-docx ``Table``, a colored/filled
callout box, a solid-color banner, or any cell/paragraph background shading.
python-docx has no high-level API for background fill; it requires injecting
low-level ``<w:shd>``/``<w:pBdr>`` OOXML elements directly. This module is
that layer, kept separate from ``docx_layout.py`` since tables/shading/oxml
are far more special-purpose than anything the plain-letter generators need —
it composes ``docx_layout``'s primitives (``add_paragraph``) rather than
duplicating them.

Built for ``interview_prep.py``'s DOCX rebuild; nothing else in this repo
uses tables or shading today, so there's no existing visual language to
match beyond ``docx_layout.BLUE`` (reused here) and the SKILL.md's own
(previously unimplemented) Step 7 color spec.
"""

from __future__ import annotations

from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from shared.docx_layout import BODY_SIZE, add_paragraph

# Hex strings (no leading '#') for oxml fill/border attributes, which take
# raw hex text, not RGBColor objects.
NAVY_HEX = "1B3A5C"
BLUE_HEX = "2E6DA4"  # matches docx_layout.BLUE
LBLUE_HEX = "D6E4F0"
AMBER_HEX = "FFF8E1"
AMBERB_HEX = "B45309"
GREEN_HEX = "E8F5E9"
GREENB_HEX = "2E7D32"
LGRAY_HEX = "F5F7FA"
MGRAY_HEX = "CCCCCC"
WHITE_HEX = "FFFFFF"
RED_HEX = "C0392B"

NAVY = RGBColor.from_string(NAVY_HEX)
LBLUE = RGBColor.from_string(LBLUE_HEX)
AMBER = RGBColor.from_string(AMBER_HEX)
AMBERB = RGBColor.from_string(AMBERB_HEX)
GREEN = RGBColor.from_string(GREEN_HEX)
GREENB = RGBColor.from_string(GREENB_HEX)
LGRAY = RGBColor.from_string(LGRAY_HEX)
MGRAY = RGBColor.from_string(MGRAY_HEX)
WHITE = RGBColor.from_string(WHITE_HEX)
RED = RGBColor.from_string(RED_HEX)


def shade_cell(cell: Any, hex_color: str) -> None:
    """Set a table cell's background fill via a raw ``<w:shd>`` element."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def shade_paragraph(paragraph: Any, hex_color: str) -> None:
    """Set a paragraph's background fill via a raw ``<w:shd>`` element."""
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def _set_paragraph_border(paragraph: Any, side: str, hex_color: str, size_eighths_pt: int) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size_eighths_pt))
    edge.set(qn("w:space"), "4")
    edge.set(qn("w:color"), hex_color)
    pBdr.append(edge)


def set_paragraph_left_border(paragraph: Any, hex_color: str, size_eighths_pt: int = 18) -> None:
    """Add a colored left border (the callout-box / card-accent idiom)."""
    _set_paragraph_border(paragraph, "left", hex_color, size_eighths_pt)


def set_paragraph_top_border(paragraph: Any, hex_color: str, size_eighths_pt: int = 6) -> None:
    """Add a colored top border (used by Q&A cards)."""
    _set_paragraph_border(paragraph, "top", hex_color, size_eighths_pt)


def two_column_table(
    doc: Any,
    headers: list[str],
    rows: list[tuple[str, str]],
    *,
    header_fill: str = NAVY_HEX,
    alt_row_fill: str = LGRAY_HEX,
) -> Any:
    """A two-column table with a colored header row and alternating row shading.

    Used for the Experience Match table (requirement -> evidence) and the
    Recruiter Briefing focus-areas table.
    """
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.autofit = True

    for col, text in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = ""
        shade_cell(cell, header_fill)
        p = cell.paragraphs[0]
        run = p.add_run(text)
        run.font.name = "Carlito"
        run.font.size = BODY_SIZE
        run.font.bold = True
        run.font.color.rgb = WHITE

    for row_idx, (left, right) in enumerate(rows, start=1):
        for col, text in enumerate((left, right)):
            cell = table.cell(row_idx, col)
            cell.text = ""
            if row_idx % 2 == 0:
                shade_cell(cell, alt_row_fill)
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.name = "Carlito"
            run.font.size = BODY_SIZE

    return table


def banner(
    doc: Any,
    *,
    title: str,
    subtitle_lines: list[str] | None = None,
    fill: str = NAVY_HEX,
    text_color: RGBColor = WHITE,
) -> Any:
    """A full-width solid-color banner (1x1 table — clean full-bleed padding
    that a shaded paragraph can't give, since paragraph shading only fills
    to the text line's box)."""
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    cell.text = ""
    shade_cell(cell, fill)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.name = "Carlito"
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = text_color

    for line in subtitle_lines or []:
        sp = cell.add_paragraph()
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        srun = sp.add_run(line)
        srun.font.name = "Carlito"
        srun.font.size = Pt(11)
        srun.font.color.rgb = text_color

    return table


def callout_box(
    doc: Any,
    label: str,
    body: str,
    *,
    fill: str = AMBER_HEX,
    border_color: str = AMBERB_HEX,
) -> Any:
    """A shaded, left-bordered callout paragraph — the "Gap note:" idiom."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    shade_paragraph(p, fill)
    set_paragraph_left_border(p, border_color)
    label_run = p.add_run(f"{label}: ")
    label_run.font.name = "Carlito"
    label_run.font.size = BODY_SIZE
    label_run.font.bold = True
    body_run = p.add_run(body)
    body_run.font.name = "Carlito"
    body_run.font.size = BODY_SIZE
    return p


def qa_card(
    doc: Any,
    question: str,
    answer: str,
    *,
    tip: str | None = None,
    gap_note: str | None = None,
) -> None:
    """A question/answer card: bold navy question with a blue top border,
    plain answer, optional amber-italic Tip line, optional gap-coaching
    callout directly below."""
    qp = doc.add_paragraph()
    qp.paragraph_format.space_before = Pt(10)
    qp.paragraph_format.space_after = Pt(2)
    set_paragraph_top_border(qp, BLUE_HEX)
    qrun = qp.add_run(question)
    qrun.font.name = "Carlito"
    qrun.font.size = BODY_SIZE
    qrun.font.bold = True
    qrun.font.color.rgb = NAVY

    add_paragraph(doc, answer, size=BODY_SIZE, space_after=Pt(2))

    if tip:
        add_paragraph(
            doc, f"Tip: {tip}", italic=True, size=Pt(9.5), color=AMBERB, space_after=Pt(2)
        )

    if gap_note:
        callout_box(doc, "Gap note", gap_note)


# (fill, text color) per row. S/R use fully-saturated fills (need white
# text for contrast); T/A use pale tints (default/dark text reads fine) —
# this asymmetry matches the SKILL.md's original color spec, which only
# defines pale-tint + darker-accent pairs for amber/green, not for blue/red.
_STAR_ROWS = (
    ("S", BLUE_HEX, WHITE),
    ("T", GREEN_HEX, None),
    ("A", AMBER_HEX, None),
    ("R", RED_HEX, WHITE),
)


def star_card(
    doc: Any,
    title: str,
    tags: list[str],
    situation: str,
    task: str,
    action: str,
    result: str,
    *,
    adapt_when: list[str] | None = None,
) -> None:
    """A STAR story card: navy header, four color-coded S/T/A/R rows, and an
    optional reference-only "Adapt when" list (never rewritten by the
    script — framing language stays an LLM judgment call, see the plan's
    scope boundary)."""
    header = doc.add_paragraph()
    header.paragraph_format.space_before = Pt(10)
    header.paragraph_format.space_after = Pt(2)
    shade_paragraph(header, NAVY_HEX)
    hrun = header.add_run(title if not tags else f"{title}  ·  {', '.join(tags)}")
    hrun.font.name = "Carlito"
    hrun.font.size = BODY_SIZE
    hrun.font.bold = True
    hrun.font.color.rgb = WHITE

    row_style = {label: (fill, text_color) for label, fill, text_color in _STAR_ROWS}
    for label, text in zip(("S", "T", "A", "R"), (situation, task, action, result), strict=True):
        fill, text_color = row_style[label]
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        shade_paragraph(p, fill)
        set_paragraph_left_border(p, fill)
        label_run = p.add_run(f"{label}: ")
        label_run.font.name = "Carlito"
        label_run.font.size = BODY_SIZE
        label_run.font.bold = True
        body_run = p.add_run(text)
        body_run.font.name = "Carlito"
        body_run.font.size = BODY_SIZE
        if text_color is not None:
            label_run.font.color.rgb = text_color
            body_run.font.color.rgb = text_color

    if adapt_when:
        add_paragraph(
            doc,
            "Adapt when: " + " · ".join(adapt_when),
            italic=True,
            size=Pt(9.5),
            color=MGRAY,
            space_before=Pt(2),
            space_after=Pt(6),
        )
